from pathlib import Path
import sqlite3
from collections import Counter
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = Path("data/calcul_distances_global/distances.sqlite")
SORTIE = Path("data/calcul_distances_global/rapport_controle_distances.xlsx")

if not BASE.exists():
    raise FileNotFoundError(f"Base introuvable : {BASE}")

con = sqlite3.connect(BASE)
con.row_factory = sqlite3.Row

try:
    lignes = con.execute("""
        SELECT
            e.employee_id,
            e.home_address,
            e.commute_mode,
            c.distance_km,
            c.status,
            c.source,
            c.detail
        FROM employee e
        JOIN commute_check c
          ON c.employee_id = e.employee_id
        ORDER BY e.employee_id
    """).fetchall()
finally:
    con.close()

sportifs = [
    r for r in lignes
    if r["status"] != "NON_APPLICABLE"
]

a_verifier = [
    r for r in sportifs
    if r["status"] == "NON_VERIFIE"
]

compteurs = Counter(r["status"] for r in lignes)

wb = Workbook()
synthese = wb.active
synthese.title = "Synthese"
trajets = wb.create_sheet("Trajets sportifs")
anomalies = wb.create_sheet("Adresses a verifier")

synthese.append(["RAPPORT DE CONTROLE DES DISTANCES"])
synthese.append(["Date de generation", datetime.now().strftime("%d/%m/%Y %H:%M")])
synthese.append(["Source", str(BASE)])
synthese.append([])
synthese.append(["Indicateur", "Nombre"])
synthese.append(["Total salaries", len(lignes)])
synthese.append(["Trajets sportifs", len(sportifs)])
synthese.append(["VALIDE", compteurs["VALIDE"]])
synthese.append(["A_VERIFIER", compteurs["A_VERIFIER"]])
synthese.append(["NON_VERIFIE", compteurs["NON_VERIFIE"]])
synthese.append(["NON_APPLICABLE", compteurs["NON_APPLICABLE"]])
synthese.append([])
synthese.append(["Note", "Les distances manquantes ne sont pas estimees."])

entetes = [
    "ID salarie",
    "Adresse domicile",
    "Mode de transport",
    "Distance (km)",
    "Statut",
    "Source",
    "Detail"
]

trajets.append(entetes)
anomalies.append(entetes)

for r in sportifs:
    valeurs = [
        r["employee_id"],
        r["home_address"],
        r["commute_mode"],
        r["distance_km"],
        r["status"],
        r["source"],
        r["detail"]
    ]
    trajets.append(valeurs)

for r in a_verifier:
    anomalies.append([
        r["employee_id"],
        r["home_address"],
        r["commute_mode"],
        r["distance_km"],
        r["status"],
        r["source"],
        r["detail"]
    ])

bleu = "17365D"
blanc = "FFFFFF"

for feuille in wb.worksheets:
    ligne_entete = 5 if feuille == synthese else 1

    for cellule in feuille[ligne_entete]:
        cellule.fill = PatternFill("solid", fgColor=bleu)
        cellule.font = Font(color=blanc, bold=True)
        cellule.alignment = Alignment(wrap_text=True)

    feuille.freeze_panes = (
        "A6" if feuille == synthese else "A2"
    )

    if feuille != synthese:
        feuille.auto_filter.ref = feuille.dimensions

    largeurs = (
        [30, 65]
        if feuille == synthese
        else [16, 48, 28, 18, 20, 30, 65]
    )

    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[
            get_column_letter(i)
        ].width = largeur

    if feuille != synthese:
        for ligne in feuille.iter_rows(min_row=2):
            ligne[3].number_format = "0.000"

synthese["A1"].font = Font(
    size=16,
    bold=True,
    color=bleu
)

SORTIE.parent.mkdir(parents=True, exist_ok=True)
wb.save(SORTIE)

print("=== RAPPORT EXCEL CREE ===")
print("Fichier :", SORTIE)
print("Salaries :", len(lignes))
print("Trajets sportifs :", len(sportifs))
print("Trajets valides :", compteurs["VALIDE"])
print("Trajets non verifies :", compteurs["NON_VERIFIE"])
print("Base du jury non modifiee.")
